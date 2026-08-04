"""Append extracted data to the existing ใบเบิกวัตถุดิบ.xlsx.

Rows 1 = header, row 2 = data-type labels, row 3+ = data (matches the file we built).
Columns are located by header name so the writer survives column reordering.
"""
import threading

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

_lock = threading.Lock()

_thin = Border(left=Side("thin"), right=Side("thin"),
               top=Side("thin"), bottom=Side("thin"))
_font = Font(name="Arial", size=11)
_ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
_lft = Alignment(horizontal="left", vertical="center", wrap_text=True)
_LEFT_COLS = {"material_name"}


def _headers(ws):
    return {ws.cell(row=1, column=c).value: c
            for c in range(1, ws.max_column + 1)
            if ws.cell(row=1, column=c).value}


def _next_row(ws):
    """First empty data row (data starts at row 3, row 2 is the type labels)."""
    r = 3
    while ws.cell(row=r, column=1).value not in (None, ""):
        r += 1
    return r


def _write_row(ws, headers, row_dict):
    r = _next_row(ws)
    for name, col in headers.items():
        cell = ws.cell(row=r, column=col, value=row_dict.get(name))
        cell.font = _font
        cell.border = _thin
        cell.alignment = _lft if name in _LEFT_COLS else _ctr
    return r


def order_exists(excel_path, production_order_no):
    """True if this production_order_no already has rows in production_actual."""
    with _lock:
        wb = load_workbook(excel_path)
        try:
            ws = wb["production_actual"]
        except KeyError:
            return False
        headers = _headers(ws)
        col = headers.get("production_order_no")
        if not col:
            return False
        for r in range(3, ws.max_row + 1):
            if str(ws.cell(row=r, column=col).value or "").strip() == str(production_order_no).strip():
                return True
        return False


def append_record(excel_path, data, source_file=None, provider=None, scanned_at=None):
    """Append one extracted form. Returns a dict summary of what was written."""
    with _lock:
        wb = load_workbook(excel_path)
        po = data.get("production_order_no")
        date = data.get("document_date")
        added = {"production_actual": 0, "production_workforce": 0, "production_pack": 0}

        # production_actual — one row per material
        if "production_actual" in wb.sheetnames:
            ws = wb["production_actual"]
            h = _headers(ws)
            for m in data.get("materials", []):
                _write_row(ws, h, {
                    "production_order_no": po,
                    "document_date": date,
                    "material_code": m.get("material_code"),
                    "material_name": m.get("material_name"),
                    "item_no": m.get("item_no"),
                    "actual_qty": m.get("actual_qty"),
                    "unit": m.get("unit"),
                })
                added["production_actual"] += 1

        # production_workforce — one row per form
        if "production_workforce" in wb.sheetnames:
            ws = wb["production_workforce"]
            h = _headers(ws)
            wf = data.get("workforce", {})
            if wf.get("worker_count") is not None or wf.get("work_hours") is not None:
                _write_row(ws, h, {
                    "production_order_no": po,
                    "document_date": date,
                    "worker_count": wf.get("worker_count"),
                    "work_hours": wf.get("work_hours"),
                })
                added["production_workforce"] += 1

        # production_pack — one row per pack line
        if "production_pack" in wb.sheetnames:
            ws = wb["production_pack"]
            h = _headers(ws)
            for p in data.get("packs", []):
                _write_row(ws, h, {
                    "production_order_no": po,
                    "document_date": date,
                    "pack_no": p.get("pack_no"),
                    "mfg_date": p.get("mfg_date"),
                    "exp_date": p.get("exp_date"),
                    "quantity": p.get("quantity"),
                    "unit": p.get("unit"),
                })
                added["production_pack"] += 1

        _append_scan_log(wb, scanned_at, source_file, provider, po, added, "success")
        wb.save(excel_path)
        return added


def log_failure(excel_path, source_file, provider, scanned_at, error):
    with _lock:
        wb = load_workbook(excel_path)
        _append_scan_log(wb, scanned_at, source_file, provider, None,
                         {}, "failed", error)
        wb.save(excel_path)


_SCAN_LOG_HEADERS = ["scanned_at", "source_file", "provider",
                     "production_order_no", "rows_added", "status", "note"]


def _append_scan_log(wb, scanned_at, source_file, provider, po, added, status, note=""):
    if "scan_log" not in wb.sheetnames:
        ws = wb.create_sheet("scan_log")
        for c, name in enumerate(_SCAN_LOG_HEADERS, 1):
            cell = ws.cell(row=1, column=c, value=name)
            cell.font = Font(name="Arial", bold=True, size=11)
            cell.border = _thin
            cell.alignment = _ctr
        widths = [20, 30, 12, 22, 26, 12, 40]
        from openpyxl.utils import get_column_letter
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    else:
        ws = wb["scan_log"]
    r = ws.max_row + 1
    while ws.cell(row=r - 1, column=1).value in (None, "") and r > 2:
        r -= 1
    rows_added = ", ".join(f"{k}:{v}" for k, v in added.items() if v) or "-"
    vals = [scanned_at, source_file, provider, po, rows_added, status, note]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = _font
        cell.border = _thin
        cell.alignment = _lft if c in (2, 7) else _ctr
