"""Unit tests for excel_export.py — build_workbook is a pure function."""
import io
import pytest
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import load_workbook
from excel_export import build_workbook, COLUMNS


def _load(orders):
    raw = build_workbook(orders)
    assert isinstance(raw, bytes)
    assert len(raw) > 0
    return load_workbook(io.BytesIO(raw))


class TestBuildWorkbook:
    def test_empty_orders(self):
        wb = _load([])
        ws = wb.active
        assert ws.title == "production_order"
        # header row only
        assert ws.cell(1, 1).value == "Production Order"
        assert ws.cell(2, 1).value is None

    def test_header_count(self):
        wb = _load([])
        ws = wb.active
        headers = [ws.cell(1, c).value for c in range(1, len(COLUMNS) + 1)]
        expected = [col[0] for col in COLUMNS]
        assert headers == expected

    def test_header_style(self):
        wb = _load([])
        ws = wb.active
        cell = ws.cell(1, 1)
        assert cell.font.bold is True
        assert cell.fill.fgColor.rgb == "00FFC000"

    def test_single_order_no_lines(self):
        orders = [{"order_no": "OD001", "lines": []}]
        wb = _load(orders)
        ws = wb.active
        # order with no lines still produces 1 data row (with empty line dict)
        assert ws.cell(2, 1).value == "OD001"

    def test_single_order_with_lines(self):
        orders = [{
            "order_no": "OD002",
            "document_date": "2026-07-01",
            "series_no": "7010101004",
            "product_name": "ขนมปัง",
            "plan_total": 100.0,
            "actual_total": 98.5,
            "plan_unit": "KG",
            "lines": [
                {"row_no": 1, "item_no": "A001", "item_description": "แป้ง",
                 "type": "Item", "quantity": 50.0, "whse": "WH1", "plan": 55.0, "unit": "KG"},
                {"row_no": 2, "item_no": "A002", "item_description": "น้ำตาล",
                 "type": "Item", "quantity": 20.0, "whse": "WH1", "plan": 22.0, "unit": "KG"},
            ],
        }]
        wb = _load(orders)
        ws = wb.active
        # 2 lines → 2 data rows
        assert ws.cell(2, 1).value == "OD002"  # row 2, col 1 = order_no
        assert ws.cell(3, 1).value == "OD002"  # repeated on each line
        assert ws.cell(2, 6).value == 1  # row_no of first line
        assert ws.cell(3, 6).value == 2  # row_no of second line
        assert ws.cell(2, 7).value == "A001"  # item_no
        assert ws.cell(3, 7).value == "A002"

    def test_header_fields_repeated_per_line(self):
        orders = [{
            "order_no": "OD003",
            "product_name": "สินค้า",
            "lines": [{"row_no": 1}, {"row_no": 2}],
        }]
        wb = _load(orders)
        ws = wb.active
        assert ws.cell(2, 1).value == "OD003"
        assert ws.cell(3, 1).value == "OD003"
        assert ws.cell(2, 4).value == "สินค้า"
        assert ws.cell(3, 4).value == "สินค้า"

    def test_multiple_orders(self):
        orders = [
            {"order_no": "A", "lines": [{"row_no": 1}]},
            {"order_no": "B", "lines": [{"row_no": 1}, {"row_no": 2}]},
        ]
        wb = _load(orders)
        ws = wb.active
        assert ws.cell(2, 1).value == "A"
        assert ws.cell(3, 1).value == "B"
        assert ws.cell(4, 1).value == "B"

    def test_number_format_applied(self):
        orders = [{"order_no": "X", "lines": [{"quantity": 3.14}]}]
        wb = _load(orders)
        ws = wb.active
        # quantity column (10th) should have number format
        assert ws.cell(2, 10).number_format == "#,##0.000"

    def test_string_values_no_number_format(self):
        orders = [{"order_no": "X", "lines": [{"quantity": "N/A"}]}]
        wb = _load(orders)
        ws = wb.active
        assert ws.cell(2, 10).value == "N/A"

    def test_freeze_panes(self):
        wb = _load([{"order_no": "X", "lines": [{}]}])
        ws = wb.active
        assert ws.freeze_panes == "A2"

    def test_auto_filter(self):
        wb = _load([{"order_no": "X", "lines": [{}]}])
        ws = wb.active
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref.startswith("A1:")


class TestBatchSheet:
    def test_no_batches_no_sheet(self):
        wb = _load([{"order_no": "X", "lines": [{}]}])
        assert "MFG_EXP" not in wb.sheetnames

    def test_batch_sheet_created(self):
        orders = [{
            "order_no": "OD100",
            "document_date": "2024-07-01",
            "series_no": "7011001001",
            "product_name": "แซนวิช",
            "plan_total": 25000,
            "actual_total": 30362,
            "lines": [{}],
            "batches": [
                {"order_no": "OD100", "mfg_date": "2024-04-01",
                 "exp_date": "2024-09-12",
                 "batch_qty": 15046, "batch_unit": "ชิ้น"},
                {"order_no": "OD100", "mfg_date": "2024-09-06",
                 "exp_date": "2025-03-06",
                 "batch_qty": 15316, "batch_unit": "ชิ้น"},
            ],
        }]
        wb = _load(orders)
        assert "MFG_EXP" in wb.sheetnames
        ws = wb["MFG_EXP"]
        assert ws.cell(1, 1).value == "Production Order"
        assert ws.cell(1, 3).value == "Item No."
        assert ws.cell(1, 5).value == "MFG Date"
        assert ws.cell(1, 6).value == "EXP Date"
        assert ws.cell(1, 7).value == "Receive Qty."
        assert ws.cell(2, 1).value == "OD100"
        assert ws.cell(2, 3).value == "7011001001"
        assert ws.cell(2, 5).value == "2024-04-01"
        assert ws.cell(2, 7).value == 15046
        assert ws.cell(3, 7).value == 15316

    def test_batch_whse_comes_from_the_product_row(self):
        orders = [{
            "order_no": "OD200",
            "product_whse": "DW-1001",
            "lines": [{"whse": "P8-PD01"}, {"whse": "P8-PD02"}],
            "batches": [{"mfg_date": "2026-08-19", "exp_date": "2026-09-17",
                         "batch_qty": 6002, "batch_unit": "ชิ้น"}],
        }]
        ws = _load(orders)["MFG_EXP"]
        assert ws.cell(1, 9).value == "คลัง"
        # the receiving warehouse, never one of the material lines' warehouses
        assert ws.cell(2, 9).value == "DW-1001"

    def test_batch_whse_says_na_when_the_form_left_it_empty(self):
        orders = [{
            "order_no": "OD201",
            "lines": [{"whse": "P8-PD02"}, {"whse": "P8-PD02"}],
            "batches": [{"mfg_date": "2026-08-12", "exp_date": "2026-08-19",
                         "batch_qty": 4386, "batch_unit": "ชิ้น"}],
        }]
        ws = _load(orders)["MFG_EXP"]
        # never borrowed from the material lines
        assert ws.cell(2, 9).value == "n/a"

    def test_batch_whse_says_na_when_blank_rather_than_missing(self):
        orders = [{
            "order_no": "OD202",
            "product_whse": "   ",
            "lines": [{}],
            "batches": [{"mfg_date": "2026-08-12", "exp_date": "2026-08-19",
                         "batch_qty": 1, "batch_unit": "ชิ้น"}],
        }]
        ws = _load(orders)["MFG_EXP"]
        assert ws.cell(2, 9).value == "n/a"

    def test_batch_sheet_multiple_orders(self):
        orders = [
            {"order_no": "A", "actual_total": 100, "lines": [{}], "batches": [
                {"order_no": "A", "mfg_date": "2024-01-01", "exp_date": "2024-06-01",
                 "batch_qty": 100, "batch_unit": "ชิ้น"},
            ]},
            {"order_no": "B", "actual_total": 200, "lines": [{}], "batches": [
                {"order_no": "B", "mfg_date": "2024-02-01", "exp_date": "2024-07-01",
                 "batch_qty": 200, "batch_unit": "ชิ้น"},
            ]},
        ]
        wb = _load(orders)
        ws = wb["MFG_EXP"]
        assert ws.cell(2, 1).value == "A"
        assert ws.cell(3, 1).value == "B"
